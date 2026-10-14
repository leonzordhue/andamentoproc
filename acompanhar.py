"""Acompanhamento processual (TJAM): monta a pasta de cada processo com PDFs oficiais e um relatório Word.

Uso:  .venv\\Scripts\\python acompanhar.py 0000000-00.0000.8.04.0000 [outro ...] [--oab 1234/AM]
      sem número: pergunta o número (ou, se rodar agendado sem terminal, lê processos.txt, um por linha)

Para cada processo cria  Desktop\\Processos TJAM\\<número>\\  com:
  - publicacoes\\AAAA-MM-DD_<tipo>_<id>.pdf   certidões oficiais do DJEN (Comunica/CNJ) de cada publicação
  - publicacoes\\AAAA-MM-DD_<tipo>_<id>.txt   texto integral da publicação
  - PROJUDI\\   pasta onde o usuário salva (Ctrl+S ou Ctrl+P > PDF) a página da consulta pública do PROJUDI,
               que tem reCAPTCHA e não é automatizada; o texto salvo entra no relatório
  - Relatorio_<número>.docx   linha do tempo com DJEN + DataJud + PROJUDI salvo + lista de PDFs
  - estado.json   o que já foi visto (o relatório marca o que é NOVO desde a última execução)
Com --oab 1234/AM, também conta as publicações do advogado no TJAM dos últimos 7 dias e salva a lista em publicacoes_OAB.xlsx.
Fontes oficiais e públicas; nenhuma proteção (captcha/anti-robô) é contornada.
"""
import argparse
import datetime as dt
import html
import json
import os
import re
import sys

import requests

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE_DIR)
from jus import alias_datajud, datajud  # noqa: E402

# Pasta de saída: JUS_SAIDA (GitHub Actions usa ./saida); no PC, Desktop\Processos TJAM.
RAIZ = os.environ.get("JUS_SAIDA") or os.path.join(os.path.expanduser("~"), "Desktop", "Processos TJAM")

# Marcos do processo em linguagem simples. Casam com o nome do movimento (TPU/CNJ no DataJud), com o tipo de
# documento das publicações do DJEN e com o texto salvo do PROJUDI. Ordem = ordem natural do rito.
MARCOS = [
    ("Distribuição", r"distribu", "O processo foi registrado e enviado a um juízo (vara/juizado)."),
    ("Citação", r"cita[çc][ãa]o|citad", "O réu (aqui, em regra, o Estado) foi chamado oficialmente para se defender."),
    ("Contestação", r"contesta", "O réu apresentou a defesa por escrito."),
    ("Réplica / impugnação", r"r[ée]plica|impugna[çc][ãa]o [àa] contesta", "O autor respondeu à defesa do réu."),
    ("Audiência", r"audi[êe]ncia", "Houve (ou foi marcada) audiência de conciliação ou instrução."),
    ("Conclusão ao juiz", r"conclus", "Os autos foram para o juiz decidir ou despachar."),
    ("Despacho", r"despacho|mero expediente", "Ordem do juiz para andamento (sem decidir o mérito)."),
    ("Decisão / tutela", r"(?<!para )decis[ãa]o|tutela|liminar", "O juiz decidiu uma questão do processo (pode ser pedido urgente)."),
    ("Sentença", r"(?<!para )(?<!cumprimento de )senten[çc]a|julgad[oa] (procedente|improcedente)|proced[êe]ncia", "O juiz julgou o pedido (procedente, parcialmente ou improcedente)."),
    ("Recurso", r"recurso|apela[çc][ãa]o|embargos|turma recursal", "Uma das partes recorreu; o caso vai (ou foi) para a Turma Recursal/Tribunal."),
    ("Trânsito em julgado", r"tr[âa]nsito em julgado|transitad", "Não cabe mais recurso: a decisão ficou definitiva."),
    ("Cumprimento / pagamento", r"cumprimento de senten[çc]a|requisi[çc][ãa]o de pequeno valor|rpv|precat[óo]rio|alvar[áa]", "Fase de cobrar e receber o valor (RPV/precatório/alvará)."),
    ("Arquivamento / baixa", r"arquiv|baixa definitiva", "O processo foi encerrado e arquivado."),
]


def marcos(eventos):
    """eventos: lista (data AAAAMMDD, texto). Devolve [(marco, primeira data, explicação)] na ordem do rito."""
    out = []
    for nome, rx, expl in MARCOS:
        datas = sorted(d for d, t in eventos if re.search(rx, t or "", re.I))
        out.append((nome, datas[0] if datas else None, expl))
    return out
COMUNICA = "https://comunicaapi.pje.jus.br/api/v1/comunicacao"
CERTIDAO = COMUNICA + "/%s/certidao"


def fmt_cnj(d):
    return "%s-%s.%s.%s.%s.%s" % (d[:7], d[7:9], d[9:13], d[13], d[14:16], d[16:])


def limpa(t):
    t = html.unescape(re.sub(r"<br\s*/?>|</p>", "\n", str(t or "")))
    return re.sub(r"[ \t]+", " ", re.sub(r"<[^>]+>", " ", t)).strip()


def comunica(params, max_paginas=20):
    """API pública do Comunica PJe (DJEN). Filtros exatos: numeroProcesso, numeroOab/ufOab, siglaTribunal, datas."""
    itens = []
    for pag in range(1, max_paginas + 1):
        r = requests.get(COMUNICA, params=dict(params, itensPorPagina=100, pagina=pag), timeout=90)
        r.raise_for_status()
        lote = r.json().get("items", [])
        itens += lote
        if len(lote) < 100:
            break
    return itens


def djen(numero):
    itens = comunica({"numeroProcesso": re.sub(r"\D", "", numero)})
    return sorted(itens, key=lambda r: r["data_disponibilizacao"])


def baixar_publicacoes(pubs, pasta):
    os.makedirs(pasta, exist_ok=True)
    saida = []
    for r in pubs:
        base = "%s_%s_%s" % (r["data_disponibilizacao"], re.sub(r"\W+", "-", str(r["tipoComunicacao"]))[:30], r["id"])
        pdf, txt = os.path.join(pasta, base + ".pdf"), os.path.join(pasta, base + ".txt")
        if not os.path.exists(pdf):
            try:
                resp = requests.get(CERTIDAO % r["hash"], timeout=60)
                if resp.ok and resp.content[:4] == b"%PDF":
                    open(pdf, "wb").write(resp.content)
            except requests.RequestException as e:
                print("  certidão não baixada (%s): %s" % (type(e).__name__, base))
        open(txt, "w", encoding="utf-8").write(limpa(r["texto"]))
        saida.append({"id": str(r["id"]), "data": r["datadisponibilizacao"], "tipo": r["tipoComunicacao"],
                      "orgao": r["nomeOrgao"], "documento": r.get("tipoDocumento"), "texto": limpa(r["texto"]),
                      "pdf": os.path.basename(pdf) if os.path.exists(pdf) else None})
    return saida


def movimentos_datajud(d):
    try:
        hits = datajud(alias_datajud(d), {"size": 10, "query": {"match": {"numeroProcesso": d}}}, tentativas=4)
    except SystemExit as e:
        return None, str(e)
    movs = []
    for h in hits:
        s = h["_source"]
        for m in s.get("movimentos", []):
            movs.append({"data": m.get("dataHora", "")[:10], "grau": s.get("grau"), "nome": m.get("nome"), "codigo": m.get("codigo")})
    return sorted(movs, key=lambda m: m["data"]), None


def ler_projudi(pasta):
    """Texto das páginas do PROJUDI que o usuário salvou (HTML, PDF ou TXT)."""
    os.makedirs(pasta, exist_ok=True)
    textos = []
    for nome in sorted(os.listdir(pasta)):
        arq = os.path.join(pasta, nome)
        if nome.lower().endswith((".html", ".htm", ".mhtml")):
            t = limpa(re.sub(r"<script.*?</script>|<style.*?</style>", " ", open(arq, encoding="latin-1", errors="ignore").read(), flags=re.S))
        elif nome.lower().endswith(".pdf"):
            import pypdf
            t = "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(arq).pages)
        elif nome.lower().endswith(".txt"):
            t = open(arq, encoding="utf-8", errors="ignore").read()
        else:
            continue
        textos.append((nome, re.sub(r"\n\s*\n+", "\n", t).strip()))
    return textos


def relatorio(numero, pasta, pubs, movs, erro_dj, projudi, novos, oab_pubs):
    from docx import Document
    from docx.shared import Pt, RGBColor
    azul = RGBColor(0x00, 0x2D, 0x72)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    def h(t, n=1):
        for r in doc.add_heading(t, level=n).runs:
            r.font.color.rgb = azul

    p = doc.add_paragraph()
    r = p.add_run("ACOMPANHAMENTO PROCESSUAL — %s (TJAM)" % numero)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(14), azul
    doc.add_paragraph("Gerado em %s por projetos/juridico/jus/acompanhar.py. Uso interno: contém nomes de partes e advogados."
                      % dt.datetime.now().strftime("%d/%m/%Y %H:%M"))
    if novos:
        h("Novidades desde a última verificação")
        for n in novos:
            doc.add_paragraph(n, style="List Bullet")

    h("Resumo em linguagem simples")
    eventos = [(pp["data"][6:] + pp["data"][3:5] + pp["data"][:2], "%s %s %s" % (pp["tipo"], pp["documento"] or "", pp["texto"][:600])) for pp in pubs]
    eventos += [(m["data"].replace("-", ""), m["nome"] or "") for m in movs or []]
    for nome_arq, t in projudi:  # sem data confiável no texto salvo: conta como evento sem data
        eventos += [("99999999", linha) for linha in t.splitlines()]
    mk = marcos(eventos)
    ocorridos = [m for m in mk if m[1]]
    if ocorridos:
        ult = ocorridos[-1]
        doc.add_paragraph("Situação atual: a última etapa identificada é %s%s. %s" % (
            ult[0].lower(), "" if ult[1] == "99999999" else " (%s/%s/%s)" % (ult[1][6:], ult[1][4:6], ult[1][:4]), ult[2]))
    tab = doc.add_table(rows=1, cols=3)
    tab.style = "Light Grid Accent 1"
    for c, t in zip(tab.rows[0].cells, ("Etapa", "Quando", "O que significa")):
        c.text = t
    for nome, data, expl in mk:
        c = tab.add_row().cells
        c[0].text = nome
        c[1].text = ("consta (data no PROJUDI)" if data == "99999999" else "%s/%s/%s" % (data[6:], data[4:6], data[:4])) if data else "ainda não consta"
        c[2].text = expl
    doc.add_paragraph("\"Ainda não consta\" significa que a etapa não aparece nas fontes consultadas até agora; pode não ter "
                      "acontecido ou ainda não ter sido publicada/indexada (ver Fontes e limites).")

    h("1. Dados do processo")
    if pubs:
        t0 = pubs[0]["texto"]
        for rot, rx in (("Classe", r"Processo: \S+ - ([^-\n]+?) - Vara"), ("Vara de origem", r"Vara Origem: (.+?) - Juiz"),
                        ("Juiz", r"Juiz: (.+?) - Data"), ("Data de vinculação", r"Data Vincula\S*: (\S+)")):
            m = re.search(rx, t0)
            if m:
                doc.add_paragraph("%s: %s" % (rot, m.group(1).strip()))
        partes = re.findall(r"(Apelante|Apelado|Autor|Réu|Requerente|Requerido|Reclamante|Reclamado)\S*: ([^\n]+?)\s*\n\s*Advogado\S*: ([^\n]+)", t0)
        for polo, nome, adv in partes:
            doc.add_paragraph("%s: %s — advogado(a): %s" % (polo, nome.strip(), adv.strip()))
    else:
        doc.add_paragraph("Nenhuma publicação no DJEN encontrada para o número.")

    h("2. Linha do tempo")
    linhas = [(pp["data"][6:] + pp["data"][3:5] + pp["data"][:2], pp["data"], "DJEN", "%s — %s (%s)" % (pp["tipo"], pp["documento"] or "", pp["orgao"])) for pp in pubs]
    for m in movs or []:
        linhas.append((m["data"].replace("-", ""), "%s/%s/%s" % (m["data"][8:10], m["data"][5:7], m["data"][:4]), "DataJud " + str(m["grau"]), "%s (%s)" % (m["nome"], m["codigo"])))
    tab = doc.add_table(rows=1, cols=3)
    tab.style = "Light Grid Accent 1"
    for c, t in zip(tab.rows[0].cells, ("Data", "Fonte", "Andamento")):
        c.text = t
    for _, data, fonte, desc in sorted(linhas):
        c = tab.add_row().cells
        c[0].text, c[1].text, c[2].text = data, fonte, desc
    doc.add_paragraph()
    if movs is None:
        doc.add_paragraph("DataJud: %s" % erro_dj)
    elif not movs:
        doc.add_paragraph("DataJud: processo ainda não indexado (processos recentes levam semanas para entrar na base do CNJ).")

    h("3. Andamentos do PROJUDI (consulta pública salva pelo usuário)")
    if projudi:
        for nome, t in projudi:
            doc.add_paragraph("Arquivo: " + nome).runs[0].bold = True
            for linha in t.splitlines()[:400]:
                if linha.strip():
                    doc.add_paragraph(linha.strip())
    else:
        doc.add_paragraph("Ainda não há página do PROJUDI salva em %s. A consulta pública do PROJUDI tem reCAPTCHA e só pode ser "
                          "feita por uma pessoa: abrir https://projudi-consulta.tjam.jus.br/publica, pesquisar o número, abrir o "
                          "processo e salvar a página (Ctrl+S) ou imprimir em PDF (Ctrl+P) nessa pasta. Na próxima execução o texto "
                          "entra aqui." % os.path.join(pasta, "PROJUDI"))

    h("4. Íntegra das publicações oficiais (DJEN)")
    for pp in pubs:
        doc.add_paragraph("%s — %s — %s" % (pp["data"], pp["tipo"], pp["orgao"])).runs[0].bold = True
        doc.add_paragraph(pp["texto"])
        doc.add_paragraph("PDF: publicacoes\\%s" % pp["pdf"] if pp["pdf"] else "PDF: certidão indisponível")

    if oab_pubs is not None:
        h("5. Publicações do advogado no TJAM (últimos 7 dias, DJEN)")
        doc.add_paragraph("%d publicação(ões) com a OAB do advogado como destinatária. A lista completa está em %s."
                          % (len(oab_pubs), os.path.join(RAIZ, "publicacoes_OAB.xlsx")))
        doc.add_paragraph("Deste processo: %d." % sum(1 for r in oab_pubs if r["numeroprocessocommascara"] == numero))

    h("Fontes e limites")
    for t in ["DJEN / Comunica PJe (CNJ): publicações e certidões oficiais em PDF.",
              "DataJud (API pública do CNJ): movimentos processuais; sem partes nem documentos; atraso de semanas para processos novos.",
              "PROJUDI/TJAM: andamento completo e documentos; consulta pública protegida por reCAPTCHA (etapa humana); documentos dos autos "
              "exigem o acesso do advogado cadastrado."]:
        doc.add_paragraph(t, style="List Bullet")
    arq = os.path.join(pasta, "Relatorio_%s.docx" % numero)
    doc.save(arq)
    return arq


def acompanhar(numero, oab):
    d = re.sub(r"\D", "", numero)
    numero = fmt_cnj(d)
    pasta = os.path.join(RAIZ, numero)
    os.makedirs(os.path.join(pasta, "PROJUDI"), exist_ok=True)
    print("== %s" % numero)
    pubs = baixar_publicacoes(djen(numero), os.path.join(pasta, "publicacoes"))
    print("  DJEN: %d publicação(ões)" % len(pubs))
    movs, erro = movimentos_datajud(d)
    print("  DataJud: %s" % (erro or "%d movimento(s)" % len(movs)))
    projudi = ler_projudi(os.path.join(pasta, "PROJUDI"))
    print("  PROJUDI salvo: %d arquivo(s)" % len(projudi))
    est_arq = os.path.join(pasta, "estado.json")
    antes = json.load(open(est_arq, encoding="utf-8")) if os.path.exists(est_arq) else None
    agora = {"djen": [p["id"] for p in pubs], "datajud": ["%s|%s|%s" % (m["data"], m["grau"], m["codigo"]) for m in movs or []],
             "projudi": [n for n, _ in projudi]}
    novos = []
    if antes:
        novos += ["DJEN: %s — %s" % (p["data"], p["tipo"]) for p in pubs if p["id"] not in antes["djen"]]
        novos += ["DataJud: %s" % m for m in agora["datajud"] if m not in antes["datajud"]]
        novos += ["PROJUDI: página salva %s" % n for n in agora["projudi"] if n not in antes["projudi"]]
    oab_pubs = None
    if oab:
        num_oab, uf = (oab.split("/") + ["AM"])[:2]
        hoje = dt.date.today()
        oab_pubs = comunica({"numeroOab": num_oab, "ufOab": uf, "siglaTribunal": "TJAM",
                             "dataDisponibilizacaoInicio": (hoje - dt.timedelta(days=7)).isoformat(),
                             "dataDisponibilizacaoFim": hoje.isoformat()})
    if oab_pubs:
        import pandas as pd
        cols = ["datadisponibilizacao", "numeroprocessocommascara", "tipoComunicacao", "nomeOrgao", "nomeClasse", "tipoDocumento"]
        pd.DataFrame([{c: r.get(c) for c in cols} | {"texto": limpa(r.get("texto"))[:3000]} for r in oab_pubs]).to_excel(
            os.path.join(RAIZ, "publicacoes_OAB.xlsx"), index=False)
    arq = relatorio(numero, pasta, pubs, movs, erro, projudi, novos, oab_pubs)
    json.dump(agora, open(est_arq, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("  novidades: %d | relatório: %s" % (len(novos), arq))
    return novos


def main():
    a = argparse.ArgumentParser()
    a.add_argument("numeros", nargs="*")
    a.add_argument("--oab", help="número/UF da OAB do advogado, ex.: 1234/AM")
    a = a.parse_args()
    lista = os.path.join(BASE_DIR, "processos.txt")
    nums = a.numeros or ([l.strip() for l in open(lista, encoding="utf-8") if l.strip() and not l.startswith("#")]
                         if os.path.exists(lista) and not sys.stdin.isatty() else [])
    if not nums:  # uso interativo: pergunta o número
        nums = [input("Número do processo (TJAM): ").strip()]
    linhas = []
    for n in nums:
        try:
            nov = acompanhar(n, a.oab)
            linhas.append("%s: %s" % (n, "; ".join(nov) if nov else "sem novidades"))
        except Exception as e:  # um processo com erro não para os outros
            linhas.append("%s: ERRO %s" % (n, e))
    os.makedirs(RAIZ, exist_ok=True)
    resumo = "Última verificação: %s\n\n%s\n" % (dt.datetime.now().strftime("%d/%m/%Y %H:%M"), "\n".join(linhas))
    open(os.path.join(RAIZ, "ULTIMA_VERIFICACAO.txt"), "w", encoding="utf-8").write(resumo)
    print(resumo)


if __name__ == "__main__":
    main()
